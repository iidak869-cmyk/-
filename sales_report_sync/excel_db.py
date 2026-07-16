# -*- coding: utf-8 -*-
"""営業報告データベースExcelの読み込み・重複チェック・追記・保存。

重複判定は要件どおり会社名のみで行う。
比較時は全角/半角・前後空白などの表記ゆれをNFKC正規化で吸収するが、
Excelへは取得した値をそのまま書き込む。
"""

import logging
import os
import unicodedata

from openpyxl import Workbook, load_workbook

from fields import FIELDS

logger = logging.getLogger(__name__)


def normalize_company_name(name: str) -> str:
    if name is None:
        return ""
    return "".join(unicodedata.normalize("NFKC", str(name)).split())


class ExcelSaveLockedError(Exception):
    """Excelが開かれている等の理由で保存できなかった場合のエラー。"""


class ExcelDatabase:
    def __init__(self, path: str, sheet_name: str = ""):
        self.path = path
        self._load(sheet_name)
        self._build_column_map()
        self._collect_existing_companies()

    def _load(self, sheet_name):
        if os.path.exists(self.path):
            self.workbook = load_workbook(self.path)
            if sheet_name and sheet_name in self.workbook.sheetnames:
                self.sheet = self.workbook[sheet_name]
            else:
                self.sheet = self.workbook.active
            logger.info("Excelを読み込みました: %s (シート: %s)", self.path, self.sheet.title)
        else:
            logger.warning("Excelが存在しないため新規作成します: %s", self.path)
            self.workbook = Workbook()
            self.sheet = self.workbook.active
            if sheet_name:
                self.sheet.title = sheet_name
            for column, field in enumerate(FIELDS, start=1):
                self.sheet.cell(row=1, column=column, value=field)

    def _build_column_map(self):
        """1行目のヘッダーから項目名→列番号の対応を作る。

        Excel側に「NG理由②」など転記対象外の列があってもそのまま残し、
        FIELDS に含まれる列だけへ書き込む。無いヘッダーは末尾に追加する。
        """
        self.column_map = {}
        max_column = self.sheet.max_column
        for column in range(1, max_column + 1):
            header = self.sheet.cell(row=1, column=column).value
            if header is None:
                continue
            key = "".join(str(header).split())
            if key in FIELDS and key not in self.column_map:
                self.column_map[key] = column

        next_column = max_column + 1 if self.sheet.cell(row=1, column=1).value is not None else 1
        for field in FIELDS:
            if field not in self.column_map:
                logger.warning("Excelにヘッダー「%s」が無いため列を追加します(列%d)。", field, next_column)
                self.sheet.cell(row=1, column=next_column, value=field)
                self.column_map[field] = next_column
                next_column += 1

    def _collect_existing_companies(self):
        self.existing_companies = set()
        name_column = self.column_map["会社名"]
        for row in range(2, self.sheet.max_row + 1):
            value = self.sheet.cell(row=row, column=name_column).value
            normalized = normalize_company_name(value)
            if normalized:
                self.existing_companies.add(normalized)
        logger.info("Excel登録済みの会社数: %d", len(self.existing_companies))

    def has_company(self, company_name: str) -> bool:
        return normalize_company_name(company_name) in self.existing_companies

    def append(self, record: dict):
        row = self.sheet.max_row + 1
        for field in FIELDS:
            self.sheet.cell(row=row, column=self.column_map[field], value=record.get(field, ""))
        self.existing_companies.add(normalize_company_name(record.get("会社名", "")))

    def save(self):
        """上書き保存する。Excelが開かれていて保存できない場合はエラーとして扱う。"""
        try:
            self.workbook.save(self.path)
        except PermissionError as error:
            raise ExcelSaveLockedError(
                f"Excelを保存できませんでした。ファイルが開かれていないか確認してください: {self.path}"
            ) from error
        logger.info("Excelを上書き保存しました: %s", self.path)
