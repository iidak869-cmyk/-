"""
テスト実行用スクリプト
ローカルのExcelファイルとモックCMSサーバーで動作確認する
"""
import asyncio
import openpyxl
from pathlib import Path
from playwright.async_api import async_playwright, Page

# テスト用設定（本番スクリプトと同じロジック）
FOLDER_PATH = Path(".")  # カレントディレクトリでテスト
TARGET_FILES = ["更新代行反映CMS1.xlsx", "更新代行反映CMS5.xlsx"]

COL_URL   = "管理画面URL"
COL_ID    = "ID"
COL_PW    = "PW"
COL_TITLE = "ブログタイトル"
COL_BODY  = "ブログ内容"


def load_excel(file_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        record = dict(zip(headers, row))
        if all(record.get(c) for c in [COL_URL, COL_ID, COL_PW, COL_TITLE, COL_BODY]):
            rows.append(record)
    wb.close()
    return rows


async def process_row(page: Page, record: dict, row_num: int):
    url   = str(record[COL_URL]).strip()
    user  = str(record[COL_ID]).strip()
    pw    = str(record[COL_PW]).strip()
    title = str(record[COL_TITLE]).strip()
    body  = str(record[COL_BODY]).strip()

    print(f"  [行{row_num}] STEP4: アクセス → {url}")
    await page.goto(url, wait_until="domcontentloaded", timeout=30000)

    # STEP4: ログイン
    id_input = page.locator('input[name="user_id"]').first
    pw_input = page.locator('input[type="password"]').first
    if await id_input.count() > 0 and await pw_input.count() > 0:
        print(f"  [行{row_num}] ログイン処理")
        await id_input.fill(user)
        await pw_input.fill(pw)
        await page.locator('button[type="submit"]').click()
        await page.wait_for_load_state("domcontentloaded")

    # STEP5: 「お知らせ」クリック
    print(f"  [行{row_num}] STEP5: 「お知らせ」クリック")
    await page.locator('a:has-text("お知らせ")').first.click()
    await page.wait_for_load_state("domcontentloaded")

    # STEP6: 「項目追加」クリック
    print(f"  [行{row_num}] STEP6: 「項目追加」クリック")
    await page.locator('a:has-text("項目追加")').first.click()
    await page.wait_for_load_state("domcontentloaded")

    # STEP7: タイトル入力
    print(f"  [行{row_num}] STEP7: タイトル入力 → 「{title}」")
    await page.locator('input[name="title"]').fill(title)

    # STEP8: 本文入力
    print(f"  [行{row_num}] STEP8: 本文入力")
    await page.locator('textarea[name="body"]').fill(body)

    # STEP9: 登録ボタンクリック
    print(f"  [行{row_num}] STEP9: 「登録」クリック")
    await page.locator('button:has-text("登録")').click()
    await asyncio.sleep(0.5)

    # STEP10: OKボタンクリック（モーダル内）
    print(f"  [行{row_num}] STEP10: 「OK」クリック")
    await page.locator('#confirmModal button:has-text("OK")').click()
    await page.wait_for_load_state("domcontentloaded")

    print(f"  [行{row_num}] 完了 ✓")


async def main():
    processed_files = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome",
        )
        context = await browser.new_context()
        page = await context.new_page()

        for filename in TARGET_FILES:
            file_path = FOLDER_PATH / filename

            # STEP1/2: ファイル存在確認
            if not file_path.exists():
                print(f"\n[スキップ] ファイルが存在しません: {filename}")
                continue

            print(f"\n{'='*50}")
            print(f"[処理開始] {filename}")
            print(f"{'='*50}")

            # STEP2-a/3: データ抽出
            rows = load_excel(file_path)
            print(f"  抽出行数: {len(rows)} 件")

            for i, record in enumerate(rows, start=1):
                try:
                    await process_row(page, record, i)
                except Exception as e:
                    print(f"  [行{i}] エラー: {e}")

            processed_files.append(file_path)

        await browser.close()

    # STEP13: ファイル削除
    print(f"\n{'='*50}")
    for fp in processed_files:
        fp.unlink()
        print(f"[削除] {fp.name}")

    print("\nすべての処理が完了しました。")


if __name__ == "__main__":
    asyncio.run(main())
