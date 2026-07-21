"""営業報告データベースExcelの読み込み・追記・保存。"""
from __future__ import annotations

import copy
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .normalize import normalize_company_name

# 列番号（1始まり）。5.4節の列構成に対応。
COL_CONTRACT_DATE = 1        # A 契約日
COL_MEETING_TIME = 2         # B 商談時間
# C, D は使用しない
COL_SALES_METHOD = 5         # E 営業方法
COL_COMPANY_NAME = 6         # F 会社名
COL_RESULT = 7               # G 結果
COL_NG_REASON_1 = 8          # H NG理由①
# I NG理由② は使用しない
COL_COMPANY_ADDRESS = 10     # J 会社所在地
COL_REPRESENTATIVE = 11      # K 代表者名
COL_TYPE = 12                # L 種類
COL_SALES_STAFF = 13         # M 営業担当
COL_ACCOMPANY = 14           # N 同行
COL_APPOINTMENT_STAFF = 15   # O アポ担当
COL_INDUSTRY = 16            # P 業種
COL_CONTRACT_TYPE = 17       # Q 契約形態
COL_AMOUNT = 18              # R 金額
COL_NET1 = 19                # S ネット1
COL_NET2 = 20                # T ネット2
COL_HP = 21                  # U HP
COL_VIDEO_SITE = 22          # V 動画サイト
COL_LP = 23                  # W LP
COL_LINE = 24                # X LINE
COL_VIDEO_CR = 25            # Y 動画CR
COL_CYTEKI = 26              # Z Cyteki
COL_CREDIT_CARD = 27         # AA クレカ所持
COL_INITIAL_COLLECTION = 28  # AB 初期回収

HEADER_ROW = 1
DELIVERY_ITEM_COLUMNS = [
    COL_NET1, COL_NET2, COL_HP, COL_VIDEO_SITE, COL_LP, COL_LINE, COL_VIDEO_CR, COL_CYTEKI,
]


class ExcelLoadError(Exception):
    """Excelを読み込めない場合のエラー（全体処理を停止する）。"""


class ExcelSaveError(Exception):
    """Excelを保存できない場合のエラー（全体処理を停止する）。"""


def load_workbook_safe(file_path: str) -> Workbook:
    path = Path(file_path)
    if not path.exists():
        raise ExcelLoadError(f"Excelファイルが見つかりません: {file_path}")
    try:
        return openpyxl.load_workbook(path)
    except InvalidFileException as e:
        raise ExcelLoadError(f"Excelファイルが破損している、または形式が不正です: {file_path} ({e})") from e
    except PermissionError as e:
        raise ExcelLoadError(
            f"Excelファイルを開けません。ファイルが開かれている可能性があります: {file_path}"
        ) from e
    except Exception as e:
        raise ExcelLoadError(f"Excelファイルの読み込みに失敗しました: {file_path} ({e})") from e


def get_sheet(wb: Workbook, sheet_name: str) -> Worksheet:
    if sheet_name not in wb.sheetnames:
        raise ExcelLoadError(f"シート「{sheet_name}」が見つかりません。存在するシート: {wb.sheetnames}")
    return wb[sheet_name]


def find_last_data_row(ws: Worksheet, company_col: int = COL_COMPANY_NAME) -> int:
    """F列（会社名）に実際に値が入っている最終行を返す。

    書式だけ設定された空白行やシートの最大行数は無視する。
    データが1件もない場合はヘッダー行番号を返す。
    """
    last_row = HEADER_ROW
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        value = ws.cell(row=row, column=company_col).value
        if value is not None and str(value).strip() != "":
            last_row = row
    return last_row


def get_existing_company_names(ws: Worksheet, last_data_row: int) -> set[str]:
    """処理開始時点でExcelに存在する会社名（正規化済み）の集合を返す。"""
    names: set[str] = set()
    for row in range(HEADER_ROW + 1, last_data_row + 1):
        value = ws.cell(row=row, column=COL_COMPANY_NAME).value
        normalized = normalize_company_name(value)
        if normalized:
            names.add(normalized)
    return names


def append_row(ws: Worksheet, next_row: int, values: dict[int, object], style_source_row: Optional[int]) -> None:
    """next_row へ値を書き込む。style_source_row が指定されていれば、その行の書式を引き継ぐ。"""
    max_col = max(values.keys())
    if style_source_row:
        for col in range(1, max_col + 1):
            src_cell = ws.cell(row=style_source_row, column=col)
            dst_cell = ws.cell(row=next_row, column=col)
            if src_cell.has_style:
                dst_cell.font = copy.copy(src_cell.font)
                dst_cell.border = copy.copy(src_cell.border)
                dst_cell.fill = copy.copy(src_cell.fill)
                dst_cell.number_format = src_cell.number_format
                dst_cell.protection = copy.copy(src_cell.protection)
                dst_cell.alignment = copy.copy(src_cell.alignment)

    for col, value in values.items():
        ws.cell(row=next_row, column=col).value = value


def save_workbook_safe(wb: Workbook, file_path: str) -> None:
    try:
        wb.save(file_path)
    except PermissionError as e:
        raise ExcelSaveError(
            f"Excelファイルへ保存できません。ファイルが開かれている可能性があります: {file_path}"
        ) from e
    except Exception as e:
        raise ExcelSaveError(f"Excelファイルの保存に失敗しました: {file_path} ({e})") from e
