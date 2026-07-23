"""
Meta広告 請求書集計スクリプト（テスト環境）

処理フロー:
  1. 指定フォルダ内のPDF請求書を1件ずつ読み取り、PDFごとにExcelファイルを作成する
  2. 作成したExcelファイルを一時保存フォルダへ保存する
  3. 一時保存フォルダのExcelファイルを1件ずつ確認し、集計用Excel
     （G:\\共有ドライブ\\RPA運用-技術課連携\\〇月請求分.xlsx の「成功」シート）へ
     以下の項目を転記する
       - キャンペーン名 / 日時 / 消化金額 / 参照番号 / クレカ番号 / アカウント名
  4. データが無い、または必要項目を正しく取得できなかったExcelファイルは
     「白紙」フォルダへ移動する

  集計用Excelがまだ存在しない月（初回実行）は、--template で指定したひな形ファイルを
  コピーして作成する。ひな形は「キャンペーン名/日時/消化金額/参照番号/クレカ番号/
  アカウント名」の6列ヘッダーを持つ「成功」シートを含んでいる必要がある。

実行環境:
  ユーザーのWindows端末（C:\\、G:\\共有ドライブへアクセスできる端末）で実行する想定。

事前準備:
  pip install -r requirements.txt

実行例:
  python aggregate_invoices.py ^
      --source "C:\\Users\\iida869\\Downloads\\202605_CS運用・TRANSCEND＆ValueCreationRoom＆株式会社Growth canvas＆Hackeer＆ピラティス" ^
      --temp-dir "C:\\Users\\iida869\\Downloads\\_一時保存" ^
      --billing-month 5 ^
      --template "G:\\共有ドライブ\\RPA運用-技術課連携\\〇月請求分_ひな形.xlsx"

  --rpa-dir と --blank-dir は既定値のままでよければ省略可（README参照）。
  対象月の集計用Excelがすでに存在する場合は --template を省略できる（そのファイルに追記する）。

実際のMeta広告請求書PDF（Facebook広告マネージャからダウンロードした支払い明細）の
テキスト構造に合わせて項目を抽出する。1PDF（1回の支払い）につき、次のような行が並ぶ:

  <アカウント名>                          ← 長い場合はPDF側で末尾が「…」で省略される
  アカウントID: <数字>
  請求書/支払日
  <YYYY/MM/DD HH:MM>                     ← 支払日時（請求書全体で1つ）
  支払い方法 支払い済み
  <カードブランド> ····<下4桁>            ← クレカ番号
  参照番号: <参照番号> ¥<合計金額>
  ...
  <キャンペーン名>
  ¥<消化金額>
  <YYYY/MM/DD H:MM>〜<YYYY/MM/DD H:MM>    ← キャンペーンの掲載期間（明細行の繰り返し単位）
  <キャンペーン名> インプレッション<N>件 ¥<消化金額>   ← 上と重複する要約行（無視する）

アカウント名はPDF側で省略されることがあるため、本スクリプトではPDFからではなく
「PDFが入っているフォルダ名」から取得する（先頭の「YYYYMM_」は除去する）。
--debug 指定でPDFから抽出した生テキストを一時保存フォルダに書き出せるので、
レイアウトが変わっていないか確認したいときはそちらを参照する。
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from dataclasses import dataclass, fields
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook

# 集計用Excel・per-PDF Excelの列順（実際の集計用Excelのヘッダー表記に合わせている）
FIELD_ORDER = ["キャンペーン名", "日時", "消化金額", "参照番号", "クレカ番号", "アカウント名"]
AGGREGATE_HEADERS = FIELD_ORDER
SUCCESS_SHEET_NAME = "成功"

REFERENCE_PATTERN = re.compile(r"参照番号[:：]\s*(\S+)")
CARD_PATTERN = re.compile(
    r"((?:Visa|Master\s*Card|JCB|AMEX|American\s*Express|Discover)\s*[·\.・]{2,}\s*\d{3,4})",
    re.IGNORECASE,
)
PAYMENT_DATETIME_LABEL = re.compile(r"請求書/支払日")
PAYMENT_DATETIME_VALUE = re.compile(r"^\d{4}/\d{1,2}/\d{1,2}\s+\d{1,2}:\d{2}$")
AMOUNT_ONLY_LINE = re.compile(r"^[¥￥][\d,]+$")
FOLDER_MONTH_PREFIX = re.compile(r"^\d{6}_")


@dataclass
class InvoiceRow:
    キャンペーン名: str = ""
    日時: str = ""
    消化金額: str = ""
    参照番号: str = ""
    クレカ番号: str = ""
    アカウント名: str = ""

    def is_complete(self) -> bool:
        return all(getattr(self, f.name).strip() for f in fields(self))

    def as_row(self) -> list[str]:
        return [getattr(self, name) for name in FIELD_ORDER]


def account_name_from_folder(source_dir: Path) -> str:
    """フォルダ名（例: 202605_CS運用・TRANSCEND＆...）先頭の年月プレフィックスを除いた名称をアカウント名とする。"""
    return FOLDER_MONTH_PREFIX.sub("", source_dir.name).strip()


def extract_header_fields(lines: list[str]) -> dict[str, str]:
    result = {"参照番号": "", "クレカ番号": "", "日時": ""}

    full_text = "\n".join(lines)

    m = REFERENCE_PATTERN.search(full_text)
    if m:
        result["参照番号"] = m.group(1).strip()

    m = CARD_PATTERN.search(full_text)
    if m:
        result["クレカ番号"] = m.group(1).strip()

    for i, line in enumerate(lines):
        if PAYMENT_DATETIME_LABEL.search(line) and i + 1 < len(lines):
            candidate = lines[i + 1].strip()
            if PAYMENT_DATETIME_VALUE.match(candidate):
                result["日時"] = candidate
            break

    return result


def extract_line_items(lines: list[str]) -> list[tuple[str, str]]:
    """(キャンペーン名, 消化金額) の一覧を抽出する。

    「¥12,958」のような金額単独行を見つけ、その直前の行をキャンペーン名とする。
    直後に続く「<キャンペーン名> インプレッション...件 ¥...」の要約行は
    金額単独行に一致しないため自然に無視される。
    """
    items: list[tuple[str, str]] = []
    for i, line in enumerate(lines):
        stripped = line.strip()
        if AMOUNT_ONLY_LINE.match(stripped) and i > 0:
            campaign = lines[i - 1].strip()
            if campaign:
                items.append((campaign, stripped))
    return items


def extract_invoice_rows(pdf_path: Path, account_name: str) -> list[InvoiceRow]:
    with pdfplumber.open(pdf_path) as pdf:
        lines: list[str] = []
        for page in pdf.pages:
            lines.extend((page.extract_text() or "").splitlines())

    header = extract_header_fields(lines)
    line_items = extract_line_items(lines)

    if not line_items:
        # 明細が拾えなくてもヘッダー情報だけの1行として返す（あとで不備判定される）
        return [InvoiceRow(アカウント名=account_name, **header)]

    rows = []
    for campaign, amount in line_items:
        rows.append(InvoiceRow(
            キャンペーン名=campaign,
            消化金額=amount,
            日時=header["日時"],
            参照番号=header["参照番号"],
            クレカ番号=header["クレカ番号"],
            アカウント名=account_name,
        ))
    return rows


def write_pdf_excel(rows: list[InvoiceRow], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "請求書データ"
    ws.append(FIELD_ORDER)
    for row in rows:
        ws.append(row.as_row())
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def read_pdf_excel(xlsx_path: Path) -> list[InvoiceRow]:
    wb = load_workbook(xlsx_path)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        if excel_row is None or all(v is None for v in excel_row):
            continue
        record = dict(zip(header, excel_row))
        rows.append(InvoiceRow(**{k: (record.get(k) or "") for k in FIELD_ORDER}))
    return rows


def load_or_create_aggregate(path: Path, template: Path | None) -> Workbook:
    if path.exists():
        return load_workbook(path)
    if template is None:
        raise FileNotFoundError(
            f"集計用Excelがまだ存在しません: {path}\n"
            f"初回はひな形ファイルを --template で指定してください。"
        )
    if not template.is_file():
        raise FileNotFoundError(f"--template で指定したひな形ファイルが見つかりません: {template}")
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, path)
    return load_workbook(path)


def get_success_sheet(wb: Workbook):
    if SUCCESS_SHEET_NAME not in wb.sheetnames:
        raise KeyError(
            f"集計用Excelに「{SUCCESS_SHEET_NAME}」シートがありません（あるシート: {wb.sheetnames}）。"
            "ひな形の構成を確認してください。"
        )
    return wb[SUCCESS_SHEET_NAME]


def next_empty_row(ws) -> int:
    """ひな形はデータの無い行にも書式だけが事前設定されているため、
    ws.max_row をそのまま使わず、実際に値が入っていない最初の行を探す。"""
    row = 2
    num_cols = len(FIELD_ORDER)
    while any(ws.cell(row=row, column=c).value is not None for c in range(1, num_cols + 1)):
        row += 1
    return row


def normalize_amount(value) -> str:
    """「￥12,958」（PDF抽出時の文字列）と 12958（Excel保存後の数値）を同一視できるようにする。"""
    return str(value).replace("￥", "").replace("¥", "").replace(",", "").strip()


def dedup_key(row: InvoiceRow) -> tuple[str, str, str]:
    """1件の請求書に複数キャンペーンの明細が含まれるため、参照番号だけでなく
    キャンペーン名・消化金額も合わせて重複判定キーとする。"""
    return (row.参照番号, row.キャンペーン名, normalize_amount(row.消化金額))


def existing_dedup_keys(ws) -> set[tuple[str, str, str]]:
    idx = {name: AGGREGATE_HEADERS.index(name) for name in ("参照番号", "キャンペーン名", "消化金額")}
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[idx["参照番号"]] is None:
            continue
        ref = str(row[idx["参照番号"]])
        campaign = str(row[idx["キャンペーン名"]])
        amount = row[idx["消化金額"]]
        keys.add((ref, campaign, normalize_amount(amount)))
    return keys


def to_cell_values(row: InvoiceRow) -> list:
    """Excel側の書式（日時は日付型、消化金額は数値）に合わせて型変換する。"""
    dt = datetime.strptime(row.日時, "%Y/%m/%d %H:%M")
    amount = int(row.消化金額.replace("￥", "").replace("¥", "").replace(",", ""))
    return [row.キャンペーン名, dt, amount, row.参照番号, row.クレカ番号, row.アカウント名]


def step1_convert_pdfs_to_excel(source_dir: Path, temp_dir: Path, debug: bool) -> list[Path]:
    pdf_paths = sorted(source_dir.glob("*.pdf"))
    if not pdf_paths:
        print(f"[警告] PDFが見つかりません: {source_dir}", file=sys.stderr)

    account_name = account_name_from_folder(source_dir)

    created: list[Path] = []
    for pdf_path in pdf_paths:
        try:
            rows = extract_invoice_rows(pdf_path, account_name)
        except Exception as e:
            print(f"[エラー] PDF読み取り失敗: {pdf_path.name} ({e})", file=sys.stderr)
            rows = [InvoiceRow()]

        out_path = temp_dir / f"{pdf_path.stem}.xlsx"
        write_pdf_excel(rows, out_path)
        created.append(out_path)
        print(f"作成: {out_path.name}")

        if debug:
            with pdfplumber.open(pdf_path) as pdf:
                raw_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            debug_path = temp_dir / f"{pdf_path.stem}.debug.txt"
            debug_path.write_text(raw_text, encoding="utf-8")

    return created


def step3_transcribe_and_sort(
    excel_paths: list[Path],
    aggregate_path: Path,
    blank_dir: Path,
    template: Path | None,
) -> tuple[int, int]:
    wb = load_or_create_aggregate(aggregate_path, template)
    ws = get_success_sheet(wb)
    known_keys = existing_dedup_keys(ws)
    next_row = next_empty_row(ws)

    blank_dir.mkdir(parents=True, exist_ok=True)
    ok_count = 0
    ng_count = 0

    for xlsx_path in excel_paths:
        rows = read_pdf_excel(xlsx_path)
        if not rows or not all(r.is_complete() for r in rows):
            dest = blank_dir / xlsx_path.name
            shutil.move(str(xlsx_path), str(dest))
            print(f"[白紙移動] 必要項目が不足: {xlsx_path.name}")
            ng_count += 1
            continue

        for row in rows:
            key = dedup_key(row)
            if key in known_keys:
                continue  # 再実行時の重複転記を防止
            for col, value in enumerate(to_cell_values(row), start=1):
                ws.cell(row=next_row, column=col, value=value)
            next_row += 1
            known_keys.add(key)

        ok_count += 1
        print(f"[転記完了] {xlsx_path.name}")

    wb.save(aggregate_path)
    return ok_count, ng_count


def build_aggregate_path(rpa_dir: Path, billing_month: int) -> Path:
    return rpa_dir / f"{billing_month}月請求分.xlsx"


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--source", required=True, type=Path, help="PDF請求書が入っているフォルダ")
    parser.add_argument("--temp-dir", required=True, type=Path, help="PDFから作成したExcelの一時保存先")
    parser.add_argument(
        "--rpa-dir",
        type=Path,
        default=Path(r"G:\共有ドライブ\RPA運用-技術課連携"),
        help="集計用Excel・白紙フォルダが置かれている共有ドライブ上のフォルダ",
    )
    parser.add_argument("--billing-month", required=True, type=int, help="対象請求月（例: 6）")
    parser.add_argument(
        "--blank-dir",
        type=Path,
        default=None,
        help="不備ファイルの移動先（省略時は <rpa-dir>/白紙）",
    )
    parser.add_argument("--debug", action="store_true", help="PDFの抽出テキストを一時保存先に書き出す")
    parser.add_argument(
        "--template",
        type=Path,
        default=None,
        help="対象月の集計用Excelが未作成の場合にコピー元とするひな形ファイル（初回のみ必要）",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)

    if not args.source.is_dir():
        print(f"[エラー] 指定フォルダが存在しません: {args.source}", file=sys.stderr)
        return 1

    args.temp_dir.mkdir(parents=True, exist_ok=True)
    blank_dir = args.blank_dir or (args.rpa_dir / "白紙")
    aggregate_path = build_aggregate_path(args.rpa_dir, args.billing_month)

    print(f"1-2. PDF → Excel変換（保存先: {args.temp_dir}）")
    excel_paths = step1_convert_pdfs_to_excel(args.source, args.temp_dir, args.debug)

    print(f"3-4. 集計用Excelへ転記（{aggregate_path}）／不備ファイルは {blank_dir} へ移動")
    try:
        ok_count, ng_count = step3_transcribe_and_sort(excel_paths, aggregate_path, blank_dir, args.template)
    except (FileNotFoundError, KeyError) as e:
        print(f"[エラー] {e}", file=sys.stderr)
        return 1

    print(f"完了: 転記{ok_count}件 / 白紙移動{ng_count}件")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
